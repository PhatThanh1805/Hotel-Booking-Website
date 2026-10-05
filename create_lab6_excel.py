import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_lab6_excel():
    wb = openpyxl.Workbook()
    
    # ----------------------------------------------------
    # STYLES DEFINITION
    # ----------------------------------------------------
    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="D9D9D9")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_sub_header = Font(name="Calibri", size=11, bold=True, color="1F497D")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=11)
    font_code = Font(name="Consolas", size=10)
    font_pass = Font(name="Calibri", size=11, bold=True, color="27AE60")
    
    fill_navy = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_soft_blue = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    fill_accent_blue = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_light_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_green_light = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_yellow_light = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thick_bottom_side = Side(border_style="medium", color="1F497D")
    double_bottom_side = Side(border_style="double", color="1F497D")
    
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_header = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)
    border_summary = Border(top=thin_border_side, bottom=double_bottom_side)

    # ----------------------------------------------------
    # SHEET 1: TỔNG QUAN & DASHBOARD
    # ----------------------------------------------------
    ws1 = wb.active
    ws1.title = "📌 Tổng Quan Dashboard"
    ws1.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "BÁO CÁO THỰC HÀNH LAB 6.1: KIỂM THỬ HỘP TRẮNG (WHITE-BOX TESTING)"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_navy
    ws1["A1"].alignment = align_center
    ws1.row_dimensions[1].height = 40
    
    ws1.merge_cells("A2:G2")
    ws1["A2"] = "Môn học: Kiểm thử phần mềm | Bài tập: Statement Coverage & Branch Coverage | Công cụ: Python OpenPyXL"
    ws1["A2"].font = font_subtitle
    ws1["A2"].fill = fill_navy
    ws1["A2"].alignment = align_center
    ws1.row_dimensions[2].height = 20
    
    # Metadata Block
    ws1["A4"] = "THÔNG TIN BÁO CÁO"
    ws1["A4"].font = font_sub_header
    
    meta_data = [
        ("Tên bài thực hành:", "Lab 6.1 - Các kỹ thuật kiểm thử hộp trắng và kiểm thử dựa trên kinh nghiệm"),
        ("Bài 1:", "Statement Testing & Statement Coverage (Độ phủ câu lệnh) - Hàm check_student_status"),
        ("Bài 2:", "Decision/Branch Testing & Branch Coverage (Độ phủ nhánh) - Hàm calculate_discount"),
        ("Bài 3:", "GV cho thêm - So sánh & Đánh giá nâng cao White-box Testing"),
        ("Mục tiêu độ phủ:", "Đạt 100% Statement Coverage & 100% Branch Coverage với số ca kiểm thử tối thiểu"),
        ("Trạng thái hoàn thành:", "PASSED - Đã đạt 100% Coverage cho toàn bộ bài tập")
    ]
    
    row_idx = 5
    for label, val in meta_data:
        ws1.cell(row=row_idx, column=1, value=label).font = font_bold
        ws1.cell(row=row_idx, column=2, value=val).font = font_pass if "PASSED" in val else font_regular
        ws1.row_dimensions[row_idx].height = 22
        row_idx += 1
        
    # Summary Table Header
    row_idx += 1
    ws1.cell(row=row_idx, column=1, value="BẢNG TỔNG HỢP KẾT QUẢ KIỂM THỬ LAB 6").font = font_sub_header
    row_idx += 1
    
    headers_s1 = ["STT", "Tên bài tập", "Kỹ thuật kiểm thử", "Số ca kiểm thử (TC)", "Đối tượng phủ", "Độ phủ đạt được", "Đánh giá"]
    for col_idx, text in enumerate(headers_s1, 1):
        cell = ws1.cell(row=row_idx, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_accent_blue
        cell.alignment = align_center
        cell.border = border_header
    ws1.row_dimensions[row_idx].height = 28
    row_idx += 1
    
    summary_rows = [
        (1, "Bài 1: check_student_status", "Statement Testing", 2, "8 / 8 Câu lệnh", "100%", "PASSED"),
        (2, "Bài 2: calculate_discount", "Branch Testing", 2, "4 / 4 Nhánh", "100%", "PASSED"),
        (3, "Bài 3: GV Cho thêm", "So sánh & Nâng cao", "N/A", "Ma trận so sánh & Mở rộng", "100%", "PASSED")
    ]
    
    for item in summary_rows:
        ws1.row_dimensions[row_idx].height = 24
        for c_i, val in enumerate(item, 1):
            cell = ws1.cell(row=row_idx, column=c_i, value=val)
            cell.font = font_pass if val == "PASSED" else font_regular
            cell.alignment = align_center if c_i in [1, 4, 6, 7] else align_left
            cell.border = border_cell
            if val == "PASSED":
                cell.fill = fill_green_light
        row_idx += 1
        
    # Explanation note in Dashboard
    row_idx += 2
    ws1.cell(row=row_idx, column=1, value="💡 GHI CHÚ PHƯƠNG PHÁP TỐI ƯU SỐ LƯỢNG CA KIỂM THỬ:").font = font_bold
    row_idx += 1
    notes = [
        "1. Bài 1 đạt 100% Statement Coverage chỉ với 2 Ca kiểm thử (score = 100 và score = 40) vì score = 100 bao phủ nhánh True của cả 2 câu lệnh if, còn score = 40 bao phủ khối else của if 1.",
        "2. Bài 2 đạt 100% Branch Coverage chỉ với 2 Ca kiểm thử (TC1: total_amount=150, is_member=True -> Phủ True/True; TC2: total_amount=50, is_member=False -> Phủ False/False).",
        "3. Chi tiết mã nguồn, phân tích dòng lệnh, bảng ma trận vết (Traceability Matrix) được trình bày đầy đủ ở các Tab tiếp theo."
    ]
    for n in notes:
        ws1.cell(row=row_idx, column=1, value=n).font = font_regular
        row_idx += 1

    # ----------------------------------------------------
    # SHEET 2: BÀI 1 - STATEMENT COVERAGE
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Bài 1 - Statement Coverage")
    ws2.views.sheetView[0].showGridLines = True
    
    # Title
    ws2.merge_cells("A1:H1")
    ws2["A1"] = "BÀI 1: STATEMENT TESTING & STATEMENT COVERAGE (KIỂM THỬ CÂU LỆNH)"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_navy
    ws2["A1"].alignment = align_center
    ws2.row_dimensions[1].height = 35
    
    # Section 1: Pseudo-code
    ws2["A3"] = "1. ĐOẠN MÃ MẪU (PSEUDO-CODE) & PHÂN TÍCH DÒNG LỆNH"
    ws2["A3"].font = font_sub_header
    
    code_lines_b1 = [
        ("Dòng 1", 'def check_student_status(score):', 'Khai báo hàm', 'Không tính (Header)'),
        ("Dòng 2", '    result_message = "Student status unknown."', 'Gán giá trị khởi tạo cho result_message', 'Câu lệnh 1 (Executable)'),
        ("Dòng 3", '    if score >= 50:', 'Kiểm tra điều kiện score >= 50', 'Câu lệnh 2 (Executable - Condition)'),
        ("Dòng 4", '        result_message = "Student passed."', 'Gán result_message khi đạt', 'Câu lệnh 3 (Executable - True Branch 1)'),
        ("Dòng 5", '    else:', 'Từ khóa cấu trúc rẽ nhánh', 'Câu lệnh 4 (Control Structure)'),
        ("Dòng 6", '        result_message = "Student failed."', 'Gán result_message khi không đạt', 'Câu lệnh 5 (Executable - False Branch 1)'),
        ("Dòng 7", '    final_status = "Evaluation complete."', 'Gán giá trị mặc định cho final_status', 'Câu lệnh 6 (Executable)'),
        ("Dòng 8", '    if score == 100:', 'Kiểm tra điều kiện score == 100', 'Câu lệnh 7 (Executable - Condition)'),
        ("Dòng 9", '        final_status = "Perfect score achieved!"', 'Gán final_status khi đạt điểm tuyệt đối', 'Câu lệnh 8 (Executable - True Branch 2)'),
        ("Dòng 10", '    return result_message + " " + final_status', 'Trả về chuỗi kết quả kết hợp', 'Câu lệnh 9 (Executable - Return)')
    ]
    
    headers_code = ["STT Dòng", "Mã Nguồn (Pseudo-code)", "Giải Thích Chức Năng", "Phân Loại Trong Statement Testing"]
    r_idx = 4
    for c_i, h_text in enumerate(headers_code, 1):
        cell = ws2.cell(row=r_idx, column=c_i, value=h_text)
        cell.font = font_header
        cell.fill = fill_accent_blue
        cell.alignment = align_center
        cell.border = border_header
    ws2.row_dimensions[r_idx].height = 25
    r_idx += 1
    
    for row_data in code_lines_b1:
        ws2.row_dimensions[r_idx].height = 22
        for c_i, val in enumerate(row_data, 1):
            cell = ws2.cell(row=r_idx, column=c_i, value=val)
            cell.font = font_code if c_i == 2 else font_regular
            cell.alignment = align_center if c_i == 1 else align_left
            cell.border = border_cell
            if "Executable" in val:
                cell.fill = fill_light_gray
        r_idx += 1

    # Section 2: Executable Statements Summary
    r_idx += 1
    ws2.cell(row=r_idx, column=1, value="2. DANH SÁCH CÁC CÂU LỆNH CÓ THỂ THỰC THI (EXECUTABLE STATEMENTS)").font = font_sub_header
    r_idx += 1
    
    exec_statements_list = [
        ("Mã Lệnh", "Nội dung câu lệnh", "Điều kiện / Bối cảnh thực thi"),
        ("S1", 'result_message = "Student status unknown."', "Luôn thực thi khi vào hàm"),
        ("S2", "if score >= 50:", "Luôn thực thi (Kiểm tra điều kiện 1)"),
        ("S3", 'result_message = "Student passed."', "Thực thi khi S2 là TRUE (score >= 50)"),
        ("S5", 'result_message = "Student failed."', "Thực thi khi S2 là FALSE (score < 50)"),
        ("S6", 'final_status = "Evaluation complete."', "Luôn thực thi sau khối if-else đầu tiên"),
        ("S7", "if score == 100:", "Luôn thực thi (Kiểm tra điều kiện 2)"),
        ("S8", 'final_status = "Perfect score achieved!"', "Thực thi khi S7 là TRUE (score == 100)"),
        ("S9", 'return result_message + " " + final_status', "Luôn thực thi trước khi thoát hàm")
    ]
    
    for idx_es, row_es in enumerate(exec_statements_list):
        ws2.row_dimensions[r_idx].height = 22
        for c_i, val in enumerate(row_es, 1):
            cell = ws2.cell(row=r_idx, column=c_i, value=val)
            if idx_es == 0:
                cell.font = font_header
                cell.fill = fill_soft_blue
                cell.alignment = align_center
                cell.border = border_header
            else:
                cell.font = font_bold if c_i == 1 else font_regular
                cell.alignment = align_center if c_i == 1 else align_left
                cell.border = border_cell
        r_idx += 1
        
    r_idx += 1
    ws2.cell(row=r_idx, column=1, value="👉 Tổng số câu lệnh có thể thực thi (Total Executable Statements): 8 câu lệnh (S1, S2, S3, S5, S6, S7, S8, S9).").font = font_bold

    # Section 3: Test Cases Design & Coverage Matrix
    r_idx += 2
    ws2.cell(row=r_idx, column=1, value="3. THIẾT KẾ BỘ CA KIỂM THỬ TỐI THIỂU ĐẠT 100% STATEMENT COVERAGE").font = font_sub_header
    r_idx += 1
    
    headers_tc_b1 = ["Mã Ca Kiểm Thử", "Đầu Vào (score)", "Kết Quả Mong Đợi (Expected Output)", "Các Câu Lệnh Được Thực Thi", "Số Lệnh Phủ", "Độ Phủ TC"]
    for c_i, h_text in enumerate(headers_tc_b1, 1):
        cell = ws2.cell(row=r_idx, column=c_i, value=h_text)
        cell.font = font_header
        cell.fill = fill_accent_blue
        cell.alignment = align_center
        cell.border = border_header
    ws2.row_dimensions[r_idx].height = 25
    r_idx += 1
    
    tc_b1_data = [
        ("Test_Case_1", "score = 100", '"Student passed. Perfect score achieved!"', "S1, S2 (True), S3, S6, S7 (True), S8, S9", "7 / 8", "87.5%"),
        ("Test_Case_2", "score = 40", '"Student failed. Evaluation complete."', "S1, S2 (False), S5, S6, S7 (False), S9", "6 / 8 (bổ sung S5)", "75.0%"),
    ]
    
    for row_tc in tc_b1_data:
        ws2.row_dimensions[r_idx].height = 24
        for c_i, val in enumerate(row_tc, 1):
            cell = ws2.cell(row=r_idx, column=c_i, value=val)
            cell.font = font_bold if c_i in [1, 2, 5, 6] else font_regular
            cell.alignment = align_center if c_i in [1, 2, 5, 6] else align_left
            cell.border = border_cell
        r_idx += 1
        
    # Statement Traceability Matrix Table
    r_idx += 1
    ws2.cell(row=r_idx, column=1, value="MA TRẬN VẾT ĐỘ PHỦ CÂU LỆNH (STATEMENT COVERAGE MATRIX)").font = font_sub_header
    r_idx += 1
    
    headers_matrix_b1 = ["Ca Kiểm Thử", "S1", "S2", "S3", "S5", "S6", "S7", "S8", "S9", "Tổng Số Lệnh Phủ"]
    for c_i, h_text in enumerate(headers_matrix_b1, 1):
        cell = ws2.cell(row=r_idx, column=c_i, value=h_text)
        cell.font = font_header
        cell.fill = fill_soft_blue
        cell.alignment = align_center
        cell.border = border_header
    ws2.row_dimensions[r_idx].height = 25
    r_idx += 1
    
    matrix_b1_rows = [
        ("Test_Case_1 (score=100)", "✓", "✓", "✓", "✗", "✓", "✓", "✓", "✓", "7 / 8"),
        ("Test_Case_2 (score=40)", "✓", "✓", "✗", "✓", "✓", "✓", "✗", "✓", "6 / 8"),
        ("BỘ CA KIỂM THỬ (TC1 + TC2)", "✓", "✓", "✓", "✓", "✓", "✓", "✓", "✓", "8 / 8 (100%)")
    ]
    
    for idx_m, r_m in enumerate(matrix_b1_rows):
        ws2.row_dimensions[r_idx].height = 24
        is_total = (idx_m == len(matrix_b1_rows) - 1)
        for c_i, val in enumerate(r_m, 1):
            cell = ws2.cell(row=r_idx, column=c_i, value=val)
            cell.font = font_bold if (is_total or c_i == 1 or c_i == 10) else font_regular
            cell.alignment = align_center if c_i > 1 else align_left
            cell.border = border_header if is_total else border_cell
            if is_total:
                cell.fill = fill_green_light
            elif val == "✓":
                cell.fill = fill_yellow_light
        r_idx += 1

    # Section 4: Coverage Calculation Formula
    r_idx += 2
    ws2.cell(row=r_idx, column=1, value="4. TÍNH TOÁN ĐỘ PHỦ CÂU LỆNH (STATEMENT COVERAGE CALCULATION)").font = font_sub_header
    r_idx += 1
    
    calc_rows_b1 = [
        ("Công thức tổng quát:", "Statement Coverage (%) = (Số lượng câu lệnh đã thực thi / Tổng số câu lệnh có thể thực thi) * 100%"),
        ("Số câu lệnh đã thực thi:", "8 câu lệnh (S1, S2, S3, S5, S6, S7, S8, S9 đã được chạy qua bởi bộ TC1 + TC2)"),
        ("Tổng số câu lệnh có thể thực thi:", "8 câu lệnh"),
        ("Kết quả độ phủ đạt được:", "= (8 / 8) * 100% = 100.0%"),
        ("Đánh giá kết quả:", "ĐẠT 100% STATEMENT COVERAGE - Không có câu lệnh chết (Dead Code) hoặc câu lệnh bị bỏ sót.")
    ]
    
    for label, val in calc_rows_b1:
        ws2.cell(row=r_idx, column=1, value=label).font = font_bold
        ws2.cell(row=r_idx, column=2, value=val).font = font_pass if "ĐẠT" in val else font_regular
        ws2.row_dimensions[r_idx].height = 22
        r_idx += 1

    # ----------------------------------------------------
    # SHEET 3: BÀI 2 - BRANCH COVERAGE
    # ----------------------------------------------------
    ws3 = wb.create_sheet(title="Bài 2 - Branch Coverage")
    ws3.views.sheetView[0].showGridLines = True
    
    # Title
    ws3.merge_cells("A1:H1")
    ws3["A1"] = "BÀI 2: DECISION / BRANCH TESTING & BRANCH COVERAGE (KIỂM THỬ NHÁNH)"
    ws3["A1"].font = font_title
    ws3["A1"].fill = fill_navy
    ws3["A1"].alignment = align_center
    ws3.row_dimensions[1].height = 35
    
    # Section 1: Pseudo-code
    ws3["A3"] = "1. ĐOẠN MÃ MẪU (PSEUDO-CODE) & PHÂN TÍCH ĐIỀU KIỆN"
    ws3["A3"].font = font_sub_header
    
    code_lines_b2 = [
        ("Dòng 1", 'def calculate_discount(total_amount, is_member):', 'Khai báo hàm với 2 tham số', 'Khởi tạo'),
        ("Dòng 2", '    discount = 0', 'Khởi tạo giá trị chiết khấu mặc định bằng 0', 'Câu lệnh 1 (Executable)'),
        ("Dòng 3", '    if total_amount > 100:', 'Điều kiện A: Kiểm tra giá trị hóa đơn > 100', 'Câu lệnh 2 (Điều kiện A)'),
        ("Dòng 4", '        discount = total_amount * 0.10', 'Nhánh A-True: Tính giảm giá 10%', 'Câu lệnh 3 (Exec - Nhánh A-True)'),
        ("Dòng 5", '    # Nhánh A-False sẽ bỏ qua Câu lệnh 3 và tiếp tục', 'Nhánh A-False: Bỏ qua giảm giá 10%', 'Nhánh Ẩn (Nhánh A-False)'),
        ("Dòng 6", '    if is_member:', 'Điều kiện B: Kiểm tra khách hàng có phải thành viên', 'Câu lệnh 4 (Điều kiện B)'),
        ("Dòng 7", '        discount = discount + 5', 'Nhánh B-True: Cộng thêm 5$ ưu đãi thành viên', 'Câu lệnh 5 (Exec - Nhánh B-True)'),
        ("Dòng 8", '    # Nhánh B-False sẽ bỏ qua Câu lệnh 5 và tiếp tục', 'Nhánh B-False: Bỏ qua ưu đãi thành viên', 'Nhánh Ẩn (Nhánh B-False)'),
        ("Dòng 9", '    return discount', 'Trả về giá trị tổng tiền giảm giá', 'Câu lệnh 6 (Executable - Return)')
    ]
    
    r_idx = 4
    for c_i, h_text in enumerate(headers_code, 1):
        cell = ws3.cell(row=r_idx, column=c_i, value=h_text)
        cell.font = font_header
        cell.fill = fill_accent_blue
        cell.alignment = align_center
        cell.border = border_header
    ws3.row_dimensions[r_idx].height = 25
    r_idx += 1
    
    for row_data in code_lines_b2:
        ws3.row_dimensions[r_idx].height = 22
        for c_i, val in enumerate(row_data, 1):
            cell = ws3.cell(row=r_idx, column=c_i, value=val)
            cell.font = font_code if c_i == 2 else font_regular
            cell.alignment = align_center if c_i == 1 else align_left
            cell.border = border_cell
            if "Nhánh" in val:
                cell.fill = fill_light_gray
        r_idx += 1

    # Section 2: List of Branches
    r_idx += 1
    ws3.cell(row=r_idx, column=1, value="2. DANH SÁCH TẤT CẢ CÁC NHÁNH (CONTROL FLOW BRANCHES)").font = font_sub_header
    r_idx += 1
    
    branches_list = [
        ("Mã Nhánh", "Tên Điểm Quyết Định", "Biểu Thức Điều Kiện", "Kết Quả Đánh Giá", "Hành Động & Câu Lệnh Thực Thi"),
        ("Branch_A_True", "Điều kiện A (Dòng 3)", "total_amount > 100", "TRUE", "Thực thi Câu lệnh 3 (discount = total_amount * 0.10)"),
        ("Branch_A_False", "Điều kiện A (Dòng 3)", "total_amount > 100", "FALSE", "Bỏ qua Câu lệnh 3, nhảy trực tiếp đến Điều kiện B (Dòng 6)"),
        ("Branch_B_True", "Điều kiện B (Dòng 6)", "is_member == True", "TRUE", "Thực thi Câu lệnh 5 (discount = discount + 5)"),
        ("Branch_B_False", "Điều kiện B (Dòng 6)", "is_member == False", "FALSE", "Bỏ qua Câu lệnh 5, nhảy trực tiếp đến lệnh Return (Dòng 9)")
    ]
    
    for idx_b, row_b in enumerate(branches_list):
        ws3.row_dimensions[r_idx].height = 22
        for c_i, val in enumerate(row_b, 1):
            cell = ws3.cell(row=r_idx, column=c_i, value=val)
            if idx_b == 0:
                cell.font = font_header
                cell.fill = fill_soft_blue
                cell.alignment = align_center
                cell.border = border_header
            else:
                cell.font = font_bold if c_i in [1, 4] else font_regular
                cell.alignment = align_center if c_i in [1, 4] else align_left
                cell.border = border_cell
        r_idx += 1
        
    r_idx += 1
    ws3.cell(row=r_idx, column=1, value="👉 Tổng số nhánh luồng điều khiển (Total Branches): 4 nhánh (A-True, A-False, B-True, B-False).").font = font_bold

    # Section 3: Test Cases Design & Branch Matrix
    r_idx += 2
    ws3.cell(row=r_idx, column=1, value="3. THIẾT KẾ BỘ CA KIỂM THỬ TỐI THIỂU ĐẠT 100% BRANCH COVERAGE").font = font_sub_header
    r_idx += 1
    
    headers_tc_b2 = ["Mã Ca Kiểm Thử", "Đầu Vào (total_amount, is_member)", "Kết Quả Mong Đợi (discount)", "Các Nhánh Được Thực Thi", "Số Nhánh Phủ", "Độ Phủ TC"]
    for c_i, h_text in enumerate(headers_tc_b2, 1):
        cell = ws3.cell(row=r_idx, column=c_i, value=h_text)
        cell.font = font_header
        cell.fill = fill_accent_blue
        cell.alignment = align_center
        cell.border = border_header
    ws3.row_dimensions[r_idx].height = 25
    r_idx += 1
    
    tc_b2_data = [
        ("Test_Case_1", "total_amount = 150, is_member = True", "20 (150*0.10 + 5 = 20)", "Branch_A_True, Branch_B_True", "2 / 4", "50.0%"),
        ("Test_Case_2", "total_amount = 50, is_member = False", "0 (Không giảm giá)", "Branch_A_False, Branch_B_False", "2 / 4", "50.0%"),
    ]
    
    for row_tc in tc_b2_data:
        ws3.row_dimensions[r_idx].height = 24
        for c_i, val in enumerate(row_tc, 1):
            cell = ws3.cell(row=r_idx, column=c_i, value=val)
            cell.font = font_bold if c_i in [1, 2, 5, 6] else font_regular
            cell.alignment = align_center if c_i in [1, 2, 5, 6] else align_left
            cell.border = border_cell
        r_idx += 1
        
    # Branch Matrix
    r_idx += 1
    ws3.cell(row=r_idx, column=1, value="MA TRẬN VẾT ĐỘ PHỦ NHÁNH (BRANCH COVERAGE MATRIX)").font = font_sub_header
    r_idx += 1
    
    headers_matrix_b2 = ["Ca Kiểm Thử", "Branch A-True", "Branch A-False", "Branch B-True", "Branch B-False", "Tổng Số Nhánh Phủ"]
    for c_i, h_text in enumerate(headers_matrix_b2, 1):
        cell = ws3.cell(row=r_idx, column=c_i, value=h_text)
        cell.font = font_header
        cell.fill = fill_soft_blue
        cell.alignment = align_center
        cell.border = border_header
    ws3.row_dimensions[r_idx].height = 25
    r_idx += 1
    
    matrix_b2_rows = [
        ("Test_Case_1 (150, True)", "✓ (True)", "✗", "✓ (True)", "✗", "2 / 4"),
        ("Test_Case_2 (50, False)", "✗", "✓ (False)", "✗", "✓ (False)", "2 / 4"),
        ("BỘ CA KIỂM THỬ (TC1 + TC2)", "✓", "✓", "✓", "✓", "4 / 4 (100%)")
    ]
    
    for idx_m, r_m in enumerate(matrix_b2_rows):
        ws3.row_dimensions[r_idx].height = 24
        is_total = (idx_m == len(matrix_b2_rows) - 1)
        for c_i, val in enumerate(r_m, 1):
            cell = ws3.cell(row=r_idx, column=c_i, value=val)
            cell.font = font_bold if (is_total or c_i in [1, 6]) else font_regular
            cell.alignment = align_center if c_i > 1 else align_left
            cell.border = border_header if is_total else border_cell
            if is_total:
                cell.fill = fill_green_light
            elif "✓" in val:
                cell.fill = fill_yellow_light
        r_idx += 1

    # Section 4: Calculation
    r_idx += 2
    ws3.cell(row=r_idx, column=1, value="4. TÍNH TOÁN ĐỘ PHỦ NHÁNH (BRANCH COVERAGE CALCULATION)").font = font_sub_header
    r_idx += 1
    
    calc_rows_b2 = [
        ("Công thức tổng quát:", "Branch Coverage (%) = (Số lượng nhánh đã thực thi / Tổng số nhánh trong luồng điều khiển) * 100%"),
        ("Số nhánh đã thực thi:", "4 nhánh (Branch A-True, A-False, B-True, B-False đã được duyệt qua đầy đủ)"),
        ("Tổng số nhánh trong hàm:", "4 nhánh (2 điều kiện if * 2 nhánh/điều kiện)"),
        ("Kết quả độ phủ đạt được:", "= (4 / 4) * 100% = 100.0%"),
        ("Đánh giá kết quả:", "ĐẠT 100% BRANCH COVERAGE - Mọi quyết định rẽ nhánh logic đều đã được kiểm chứng.")
    ]
    
    for label, val in calc_rows_b2:
        ws3.cell(row=r_idx, column=1, value=label).font = font_bold
        ws3.cell(row=r_idx, column=2, value=val).font = font_pass if "ĐẠT" in val else font_regular
        ws3.row_dimensions[r_idx].height = 22
        r_idx += 1

    # ----------------------------------------------------
    # SHEET 4: BÀI 3 - SO SÁNH & NÂNG CAO (GV CHO THÊM)
    # ----------------------------------------------------
    ws4 = wb.create_sheet(title="Bài 3 - So Sánh & Mở Rộng")
    ws4.views.sheetView[0].showGridLines = True
    
    # Title
    ws4.merge_cells("A1:G1")
    ws4["A1"] = "BÀI 3: PHẦN NÂNG CAO - SO SÁNH STATEMENT COVERAGE VÀ BRANCH COVERAGE"
    ws4["A1"].font = font_title
    ws4["A1"].fill = fill_navy
    ws4["A1"].alignment = align_center
    ws4.row_dimensions[1].height = 35
    
    ws4["A3"] = "1. BẢNG SO SÁNH CHI TIẾT GIỮA STATEMENT COVERAGE VÀ BRANCH COVERAGE"
    ws4["A3"].font = font_sub_header
    
    headers_comp = ["Tiêu Chí So Sánh", "Statement Coverage (Phủ Câu Lệnh)", "Branch Coverage (Phủ Nhánh)", "Nhận Xét & Đánh Giá Nâng Cao"]
    r_idx = 4
    for c_i, h_text in enumerate(headers_comp, 1):
        cell = ws4.cell(row=r_idx, column=c_i, value=h_text)
        cell.font = font_header
        cell.fill = fill_accent_blue
        cell.alignment = align_center
        cell.border = border_header
    ws4.row_dimensions[r_idx].height = 25
    r_idx += 1
    
    comparison_matrix = [
        ("Định nghĩa", "Đo lường tỉ lệ phần trăm các câu lệnh thực thi đã được chạy ít nhất 1 lần.", "Đo lường tỉ lệ phần trăm các nhánh (True/False) của mọi điểm quyết định đã được kiểm thử.", "Branch Coverage kiểm soát luồng điều khiển sâu hơn Statement Coverage."),
        ("Mục tiêu chính", "Đảm bảo không có câu lệnh nào bị lãng phí (dead code) hoặc chưa được thực thi.", "Đảm bảo tất cả các kịch bản rẽ nhánh logic (True và False) đều được xử lý đúng.", "Branch Coverage giúp phát hiện lỗi thiếu xử lý khối else."),
        ("Độ mạnh kiểm thử", "Yếu hơn (Weak coverage criteria). Đạt 100% Statement chưa chắc đạt 100% Branch.", "Mạnh hơn (Stronger coverage criteria). 100% Branch Coverage LUÔN ĐẢM BẢO 100% Statement Coverage.", "Branch Coverage bao hàm toàn bộ Statement Coverage (ngoại trừ hàm không có nhánh nào)."),
        ("Số ca kiểm thử yêu cầu", "Thường ít hơn hoặc bằng số ca kiểm thử của Branch Coverage.", "Thường nhiều hơn hoặc bằng số ca kiểm thử của Statement Coverage.", "Với các câu lệnh if không có else, Branch Coverage bắt buộc phải có test case phủ nhánh False."),
        ("Phát hiện lỗi logic", "Không phát hiện được lỗi khi điều kiện sai mà không có khối else xử lý.", "Phát hiện tốt các lỗi ẩn do thiếu nhánh xử lý ngoại lệ.", "Branch Coverage an toàn hơn trong các hệ thống đòi hỏi độ tin cậy cao.")
    ]
    
    for r_comp in comparison_matrix:
        ws4.row_dimensions[r_idx].height = 28
        for c_i, val in enumerate(r_comp, 1):
            cell = ws4.cell(row=r_idx, column=c_i, value=val)
            cell.font = font_bold if c_i == 1 else font_regular
            cell.alignment = align_center if c_i == 1 else align_left
            cell.border = border_cell
        r_idx += 1
        
    # Section 2: Illustrative Example of Difference
    r_idx += 2
    ws4.cell(row=r_idx, column=1, value="2. VÍ DỤ MINH HỌA SỰ KHÁC BIỆT (STATEMENT COVERAGE VS BRANCH COVERAGE)").font = font_sub_header
    r_idx += 1
    
    example_text = [
        ("Đoạn mã ví dụ:", "if (x > 0) { print('Dương'); }  // Không có khối else"),
        ("Với 1 Test Case (x = 5):", "- Đã thực thi lệnh print('Dương') -> Đạt 100% STATEMENT COVERAGE."),
        ("", "- Tuy nhiên, nhánh (x > 0) = FALSE chưa bao giờ được kiểm thử -> CHỈ ĐẠT 50% BRANCH COVERAGE."),
        ("Kết luận rút ra:", "Muốn đạt 100% Branch Coverage, bắt buộc phải bổ sung Test Case (x = -1 hoặc x = 0) để phủ nhánh False!"),
        ("Ứng dụng thực tế:", "Trong các dự án kiểm thử phần mềm thực tế, doanh nghiệp luôn ưu tiên tiêu chuẩn Branch Coverage hoặc MCDC (Modified Condition/Decision Coverage) để đảm bảo chất lượng phần mềm tốt nhất.")
    ]
    
    for label, val in example_text:
        ws4.cell(row=r_idx, column=1, value=label).font = font_bold
        ws4.cell(row=r_idx, column=2, value=val).font = font_regular
        ws4.row_dimensions[r_idx].height = 22
        r_idx += 1

    # ----------------------------------------------------
    # AUTO ADJUST COLUMN WIDTHS FOR ALL SHEETS
    # ----------------------------------------------------
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                # Avoid title line skewing column width calculation
                if cell.row in [1, 2]:
                    continue
                # Simple width heuristic
                lines = val_str.split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    # Set explicit custom column widths for key columns for perfect visual layout
    ws1.column_dimensions['A'].width = 25
    ws1.column_dimensions['B'].width = 40
    ws1.column_dimensions['C'].width = 25
    ws1.column_dimensions['D'].width = 20
    ws1.column_dimensions['E'].width = 25
    ws1.column_dimensions['F'].width = 18
    ws1.column_dimensions['G'].width = 15

    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 42
    ws2.column_dimensions['C'].width = 45
    ws2.column_dimensions['D'].width = 42
    ws2.column_dimensions['E'].width = 15
    ws2.column_dimensions['F'].width = 15

    ws3.column_dimensions['A'].width = 22
    ws3.column_dimensions['B'].width = 42
    ws3.column_dimensions['C'].width = 45
    ws3.column_dimensions['D'].width = 42
    ws3.column_dimensions['E'].width = 15
    ws3.column_dimensions['F'].width = 15

    ws4.column_dimensions['A'].width = 25
    ws4.column_dimensions['B'].width = 45
    ws4.column_dimensions['C'].width = 45
    ws4.column_dimensions['D'].width = 45

    output_path = "e:/School/KiemThu/Hotel-Booking-Website/Lab6_KiemThuHopTrang_BaoCao.xlsx"
    wb.save(output_path)
    print(f"Successfully generated Lab 6 Excel report at: {output_path}")

if __name__ == "__main__":
    create_lab6_excel()
