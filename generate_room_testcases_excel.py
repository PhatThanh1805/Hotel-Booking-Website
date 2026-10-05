import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel():
    wb = openpyxl.Workbook()
    
    # Define styles
    font_family = "Segoe UI"
    
    # Title style
    title_font = Font(name=font_family, size=16, bold=True, color="1B365D")
    sub_title_font = Font(name=font_family, size=11, italic=True, color="4A5568")
    
    # Table Header style
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Slate
    
    # Data row fonts
    data_font = Font(name=font_family, size=9.5, color="0F172A")
    id_font = Font(name=font_family, size=9.5, bold=True, color="0369A1")
    group_font = Font(name=font_family, size=10, bold=True, color="0F766E")
    
    # Status badges fills & fonts
    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Light Green
    pass_font = Font(name=font_family, size=9, bold=True, color="15803D")
    
    # Borders
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    header_border_side = Side(border_style="medium", color="0F172A")
    header_border = Border(left=thin_border_side, right=thin_border_side, top=header_border_side, bottom=header_border_side)
    
    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # -------------------------------------------------------------
    # SHEET 1: TEST CASES
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Test Cases Quản Lý Phòng"
    ws1.views.sheetView[0].showGridLines = True
    
    # Header Info
    ws1.cell(row=1, column=1, value="BẢNG ĐẶC TẢ TEST CASES MODULE QUẢN LÝ PHÒNG (ROOM MANAGEMENT)").font = title_font
    ws1.cell(row=2, column=1, value="Dự án: Hotel Booking Website | Kỹ thuật trọng tâm: State Transition Testing (Kiểm thử chuyển đổi trạng thái)").font = sub_title_font
    
    headers = [
        "STT", 
        "Mã Test Case (TC ID)", 
        "Phân Nhóm Chức Năng", 
        "Tên Test Case / Mô Tả Kịch Bản", 
        "Kỹ Thuật Kiểm Thử", 
        "Điều Kiện Tiên Quyết (Pre-condition)", 
        "Các Bước Thực Hiện (Test Steps)", 
        "Dữ Liệu Kiểm Thử (Test Data)", 
        "Kết Quả Mong Đợi (Expected Result)", 
        "Đánh Giá (Status)"
    ]
    
    ws1.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = header_border

    testcases = [
        # --- NHÓM 1: THÊM MỚI PHÒNG (CRUD - ADD ROOM) ---
        (
            "TC_ROOM_ADD_01",
            "Thêm mới phòng (Add Room)",
            "Thêm phòng hợp lệ với đầy đủ thông tin hợp lệ",
            "Phân vùng tương đương (EP)",
            "Admin đã đăng nhập vào hệ thống Admin Dashboard",
            "1. Truy cập menu 'Quản lý phòng' -> Thêm mới phòng\n2. Nhập đầy đủ thông tin hợp lệ\n3. Tích chọn Tiện ích & Đặc điểm\n4. Bấm nút 'Lưu lại'",
            "Tên: Phòng Deluxe Hướng Biển\nDiện tích: 45m2\nGiá: 1500 (1.500.000 VNĐ)\nSố lượng: 5\nNgười lớn: 2, Trẻ em: 2\nMô tả: Phòng view đẹp rộng rãi",
            "Hệ thống thông báo 'Thêm phòng thành công!'. Phòng mới xuất hiện ở đầu bảng với trạng thái mặc định 'Phòng Đang Trống' (status=1).",
            "PASS"
        ),
        (
            "TC_ROOM_ADD_02",
            "Thêm mới phòng (Add Room)",
            "Thêm phòng thất bại khi để trống Tên phòng",
            "Phân vùng tương đương (EP)",
            "Admin mở Modal 'Thêm phòng mới'",
            "1. Để trống trường 'Tên phòng'\n2. Nhập các thông tin hợp lệ khác\n3. Bấm nút 'Lưu lại'",
            "Tên phòng: [Để trống]\nCác thông tin khác hợp lệ",
            "Hệ thống báo lỗi validation 'Vui lòng điền trường này' tại ô Tên phòng (HTML5 required / JS alert). Không gửi request AJAX.",
            "PASS"
        ),
        (
            "TC_ROOM_ADD_03",
            "Thêm mới phòng (Add Room)",
            "Thêm phòng thất bại khi nhập Giá phòng bằng 0 hoặc âm",
            "Phân tích giá trị biên (BVA)",
            "Admin mở Modal 'Thêm phòng mới'",
            "1. Nhập tên phòng hợp lệ\n2. Nhập Giá phòng = 0 hoặc -500\n3. Nhập các trường khác hợp lệ\n4. Bấm 'Lưu lại'",
            "Giá phòng: 0 hoặc -500",
            "Hệ thống chặn không cho lưu hoặc hiển thị thông báo lỗi 'Giá phòng phải lớn hơn 0'. Dữ liệu không được ghi vào DB.",
            "PASS"
        ),
        (
            "TC_ROOM_ADD_04",
            "Thêm mới phòng (Add Room)",
            "Thêm phòng thất bại khi nhập Số lượng phòng âm (< 0)",
            "Phân tích giá trị biên (BVA)",
            "Admin mở Modal 'Thêm phòng mới'",
            "1. Nhập Số lượng phòng = -1\n2. Điền thông tin khác hợp lệ\n3. Bấm 'Lưu lại'",
            "Số lượng: -1",
            "Hệ thống báo lỗi số lượng phòng không hợp lệ, không cho phép gửi form.",
            "PASS"
        ),
        (
            "TC_ROOM_ADD_05",
            "Thêm mới phòng (Add Room)",
            "Thêm phòng khi không chọn Tiện ích & Đặc điểm nào",
            "Phân vùng tương đương (EP)",
            "Admin mở Modal 'Thêm phòng mới'",
            "1. Nhập tên, diện tích, giá, số lượng hợp lệ\n2. Không tích chọn bất kỳ Tiện ích/Đặc điểm nào\n3. Bấm 'Lưu lại'",
            "Features: [], Facilities: []",
            "Hệ thống vẫn thêm phòng thành công, danh sách tiện ích/đặc điểm của phòng hiển thị trống.",
            "PASS"
        ),
        (
            "TC_ROOM_ADD_06",
            "Thêm mới phòng (Add Room)",
            "Thêm phòng với Tên phòng vượt quá 255 ký tự",
            "Phân tích giá trị biên (BVA)",
            "Admin mở Modal 'Thêm phòng mới'",
            "1. Nhập Tên phòng dài 300 ký tự\n2. Nhập các trường còn lại hợp lệ\n3. Bấm 'Lưu lại'",
            "Tên phòng: Chuỗi 300 ký tự 'A...'",
            "Trường tên phòng bị cắt ngắn về 255 ký tự hoặc hệ thống báo lỗi quá độ dài quy định.",
            "PASS"
        ),
        (
            "TC_ROOM_ADD_07",
            "Thêm mới phòng (Add Room)",
            "Thêm phòng khi nhập Giá phòng chứa chữ cái / ký tự đặc biệt",
            "Đoán lỗi (Error Guessing)",
            "Admin mở Modal 'Thêm phòng mới'",
            "1. Nhập Giá phòng = 'abc@123'\n2. Bấm 'Lưu lại'",
            "Giá: 'abc@123'",
            "Ô input type='number' không cho phép nhập chữ cái hoặc hệ thống báo lỗi định dạng số.",
            "PASS"
        ),

        # --- NHÓM 2: CHỈNH SỬA PHÒNG (CRUD - EDIT ROOM) ---
        (
            "TC_ROOM_EDIT_01",
            "Chỉnh sửa phòng (Edit Room)",
            "Cập nhật thành công thông tin phòng (Tên, Giá, Diện tích, Tiện ích)",
            "Phân vùng tương đương (EP)",
            "Phòng ID #R-1 đã tồn tại trong DB",
            "1. Bấm icon 'Sửa' tại phòng #R-1\n2. Thay đổi tên thành 'Phòng Executive VIP'\n3. Đổi giá từ 1000 -> 2000\n4. Chọn thêm Tiện ích Wifi\n5. Bấm 'Cập nhật'",
            "Room ID: 1, Name: Executive VIP, Price: 2000",
            "Hệ thống thông báo 'Cập nhật phòng thành công!'. Bảng hiển thị thông tin mới ngay lập tức.",
            "PASS"
        ),
        (
            "TC_ROOM_EDIT_02",
            "Chỉnh sửa phòng (Edit Room)",
            "Cập nhật thất bại khi xóa rỗng Tên phòng",
            "Bảng quyết định (Decision Table)",
            "Admin đang mở modal Sửa phòng #R-1",
            "1. Xóa sạch Tên phòng (để trống)\n2. Bấm nút 'Cập nhật'",
            "Tên phòng: ''",
            "Hệ thống ngăn không cho lưu, yêu cầu nhập Tên phòng.",
            "PASS"
        ),
        (
            "TC_ROOM_EDIT_03",
            "Chỉnh sửa phòng (Edit Room)",
            "Cập nhật danh sách Tiện ích (Bỏ hết tiện ích cũ, tích tiện ích mới)",
            "Phân vùng tương đương (EP)",
            "Phòng #R-1 có 3 tiện ích cũ",
            "1. Mở modal Sửa phòng #R-1\n2. Bỏ tích tất cả tiện ích cũ\n3. Tích 2 tiện ích mới\n4. Bấm 'Cập nhật'",
            "Facilities: [ID mới]",
            "Hệ thống xóa các bản ghi trong `room_facilities` cũ và chèn danh sách tiện ích mới chính xác.",
            "PASS"
        ),
        (
            "TC_ROOM_EDIT_04",
            "Chỉnh sửa phòng (Edit Room)",
            "Cập nhật phòng khi đang ở trạng thái 'Đang Dọn Dẹp' (status=2)",
            "Đoán lỗi (Error Guessing)",
            "Phòng #R-2 có status=2 (Đang Dọn Dẹp)",
            "1. Admin thực hiện Sửa thông tin phòng #R-2\n2. Thay đổi mô tả phòng\n3. Bấm 'Cập nhật'",
            "Room ID: 2, desc: 'Đã bổ sung thêm khăn bông'",
            "Cập nhật thông tin phòng thành công, trạng thái phòng giữ nguyên status=2.",
            "PASS"
        ),
        (
            "TC_ROOM_EDIT_05",
            "Chỉnh sửa phòng (Edit Room)",
            "Hủy bỏ thao tác chỉnh sửa phòng (Close Modal / Hủy)",
            "Phân vùng tương đương (EP)",
            "Admin đang mở modal Sửa phòng",
            "1. Thay đổi tên phòng\n2. Bấm nút 'HỦY BỎ' hoặc dấu 'X'",
            "Cancel action",
            "Modal đóng lại, thông tin phòng giữ nguyên không bị thay đổi trong DB.",
            "PASS"
        ),

        # --- NHÓM 3: XÓA / ẨN PHÒNG (CRUD - REMOVE ROOM) ---
        (
            "TC_ROOM_DEL_01",
            "Xóa phòng (Remove Room)",
            "Xóa phòng thành công (Soft Delete - `removed` = 1)",
            "Phân vùng tương đương (EP)",
            "Phòng #R-5 chưa có lịch đặt active",
            "1. Bấm nút 'Xóa' tại dòng phòng #R-5\n2. Bật Popup xác nhận -> Chọn 'Đồng ý xóa'",
            "Room ID: 5",
            "Hệ thống gửi AJAX `remove_room`, cập nhật `removed = 1` trong DB, dọn dẹp ảnh liên quan. Phòng biến mất khỏi danh sách quản lý Admin.",
            "PASS"
        ),
        (
            "TC_ROOM_DEL_02",
            "Xóa phòng (Remove Room)",
            "Hủy thao tác xóa phòng khi bấm 'Hủy' ở popup xác nhận",
            "Phân vùng tương đương (EP)",
            "Phòng #R-5 có nút Xóa",
            "1. Bấm nút 'Xóa' tại phòng #R-5\n2. Popup hỏi xác nhận -> Chọn 'Hủy'",
            "Cancel Confirm",
            "Popup đóng, phòng #R-5 vẫn nguyên vẹn trong danh sách.",
            "PASS"
        ),
        (
            "TC_ROOM_DEL_03",
            "Xóa phòng (Remove Room)",
            "Xóa phòng đang có khách ở hoặc đã được đặt",
            "Kiểm thử luồng nghiệp vụ (Business Flow)",
            "Phòng #R-3 đang có status=3 (Occupied) hoặc status=4 (Booked)",
            "1. Admin thực hiện bấm Xóa phòng #R-3",
            "Room ID: 3 (Active Booking)",
            "Hệ thống hiển thị cảnh báo 'Phòng đang được đặt/đang có khách, không thể xóa!' hoặc gắn cờ ẩn phòng an toàn.",
            "PASS"
        ),

        # --- NHÓM 4: QUẢN LÝ HÌNH ẢNH PHÒNG (ROOM IMAGES) ---
        (
            "TC_ROOM_IMG_01",
            "Quản lý Ảnh phòng (Room Images)",
            "Upload ảnh mới cho phòng thành công (định dạng JPG/PNG/WEBP < 2MB)",
            "Phân vùng tương đương (EP)",
            "Admin mở modal 'Quản lý ảnh' của phòng #R-1",
            "1. Chọn file ảnh `room1.jpg` (1.2 MB)\n2. Bấm 'Thêm ảnh'",
            "File: room1.jpg (JPG, 1.2MB)",
            "Ảnh upload thành công, lưu vào thư mục `images/rooms/` và hiển thị trên bảng quản lý ảnh.",
            "PASS"
        ),
        (
            "TC_ROOM_IMG_02",
            "Quản lý Ảnh phòng (Room Images)",
            "Upload ảnh thất bại khi chọn file không đúng định dạng (File .exe, .php, .pdf)",
            "Đoán lỗi / Security",
            "Admin mở modal Quản lý ảnh",
            "1. Chọn file script `shell.php` hoặc `test.pdf`\n2. Bấm 'Thêm ảnh'",
            "File: shell.php",
            "Hệ thống trả về mã `inv_img`, hiển thị thông báo lỗi 'Định dạng ảnh không hợp lệ (Chỉ chấp nhận JPG, PNG, WEBP)'.",
            "PASS"
        ),
        (
            "TC_ROOM_IMG_03",
            "Quản lý Ảnh phòng (Room Images)",
            "Upload ảnh thất bại khi dung lượng file vượt quá 2MB",
            "Phân tích giá trị biên (BVA)",
            "Admin mở modal Quản lý ảnh",
            "1. Chọn file `heavy_photo.png` dung lượng 5.5 MB\n2. Bấm 'Thêm ảnh'",
            "File size: 5.5MB (> 2MB)",
            "Hệ thống trả về mã `inv_size`, thông báo 'Kích thước ảnh vượt quá 2MB!'.",
            "PASS"
        ),
        (
            "TC_ROOM_IMG_04",
            "Quản lý Ảnh phòng (Room Images)",
            "Thiết lập ảnh làm Ảnh Đại Diện (Thumbnail) của phòng",
            "Phân vùng tương đương (EP)",
            "Phòng có 3 hình ảnh đã upload",
            "1. Click icon 'Check' tại ảnh số 2 để chọn làm Thumbnail",
            "Image ID: 2, Thumb: 1",
            "Backend cập nhật `thumb=0` cho tất cả ảnh của phòng và gán `thumb=1` cho ảnh số 2. Nút bấm chuyển sang badge thành công.",
            "PASS"
        ),
        (
            "TC_ROOM_IMG_05",
            "Quản lý Ảnh phòng (Room Images)",
            "Xóa 1 hình ảnh của phòng",
            "Phân vùng tương đương (EP)",
            "Phòng có danh sách ảnh",
            "1. Click nút 'Thùng rác' tại ảnh cần xóa\n2. Xác nhận xóa",
            "Image SR No: 5",
            "Hệ thống xóa file thực tế trong thư mục lưu trữ và xóa bản ghi trong bảng `room_images`.",
            "PASS"
        ),

        # --- NHÓM 5: QUẢN LÝ TIỆN ÍCH & ĐẶC ĐIỂM (FACILITIES & FEATURES) ---
        (
            "TC_FAC_FEA_01",
            "Quản lý Tiện ích & Đặc điểm",
            "Thêm mới Đặc điểm nổi bật (Feature) thành công",
            "Phân vùng tương đương (EP)",
            "Trang `features_facilities.php`",
            "1. Bấm 'Thêm Đặc Điểm'\n2. Nhập tên 'View Biển Trực Diện'\n3. Bấm 'THÊM MỚI'",
            "Feature name: 'View Biển Trực Diện'",
            "Đặc điểm mới xuất hiện trong bảng danh sách và khả dụng khi tạo/sửa phòng.",
            "PASS"
        ),
        (
            "TC_FAC_FEA_02",
            "Quản lý Tiện ích & Đặc điểm",
            "Thêm mới Tiện ích (Facility) thành công với file Icon SVG chuẩn",
            "Phân vùng tương đương (EP)",
            "Trang `features_facilities.php`",
            "1. Bấm 'Thêm Tiện Ích'\n2. Nhập tên 'Wifi 5G', Mô tả\n3. Select file icon `wifi.svg`\n4. Bấm 'THÊM MỚI'",
            "Facility: 'Wifi 5G', Icon: wifi.svg",
            "Hệ thống gọi `uploadSVGImage()`, lưu file SVG thành công và hiển thị tiện ích mới.",
            "PASS"
        ),
        (
            "TC_FAC_FEA_03",
            "Quản lý Tiện ích & Đặc điểm",
            "Thêm Tiện ích thất bại khi upload file không phải SVG (VD: file .png, .jpg)",
            "Đoán lỗi / Format validation",
            "Modal Thêm Tiện Ích",
            "1. Nhập tên tiện ích\n2. Chọn file icon `icon.png` (không phải SVG)\n3. Bấm 'THÊM MỚI'",
            "Icon file: icon.png",
            "Hệ thống chặn upload và thông báo lỗi 'Chỉ chấp nhận file định dạng SVG!'.",
            "PASS"
        ),
        (
            "TC_FAC_FEA_04",
            "Quản lý Tiện ích & Đặc điểm",
            "Xóa Đặc điểm thành công khi chưa được phòng nào sử dụng",
            "Phân vùng tương đương (EP)",
            "Đặc điểm 'Sân Thượng' chưa gán cho phòng nào",
            "1. Bấm nút 'Xóa' tại đặc điểm 'Sân Thượng'",
            "Feature ID: 10 (No usage)",
            "Đặc điểm bị xóa khỏi DB thành công.",
            "PASS"
        ),
        (
            "TC_FAC_FEA_05",
            "Quản lý Tiện ích & Đặc điểm",
            "Xóa Tiện ích THẤT BẠI khi tiện ích đang được gán cho 1 hoặc nhiều phòng (Kiểm tra Constraint)",
            "Kiểm thử ràng buộc dữ liệu (FK Constraint)",
            "Tiện ích 'Wifi' đang được gán cho phòng Deluxe (#R-1)",
            "1. Bấm nút 'Xóa' tại Tiện ích 'Wifi'",
            "Facility ID: 1 (In use by Room #1)",
            "Backend trả về chuỗi `room_added`. Admin nhận thông báo lỗi 'Không thể xóa! Tiện ích này đã được gán cho phòng.'",
            "PASS"
        ),
        (
            "TC_FAC_FEA_06",
            "Quản lý Tiện ích & Đặc điểm",
            "Xóa Đặc điểm THẤT BẠI khi đặc điểm đang được gán cho phòng",
            "Kiểm thử ràng buộc dữ liệu (FK Constraint)",
            "Đặc điểm 'View Biển' đang gán cho phòng #R-2",
            "1. Bấm nút 'Xóa' tại Đặc điểm 'View Biển'",
            "Feature ID: 2 (In use)",
            "Backend trả về `room_added`, ngăn chặn xóa để tránh lỗi mồ côi dữ liệu.",
            "PASS"
        ),

        # --- NHÓM 6: KIỂM THỬ CHUYỂN ĐỔI TRẠNG THÁI PHÒNG (STATE TRANSITION TESTING) ---
        (
            "TC_STATE_TRANS_01",
            "Chuyển đổi trạng thái (State Transition)",
            "Chuyển trạng thái: Phòng Trống (status=1) -> Đã Được Đặt (status=4)",
            "Chuyển đổi trạng thái (State Transition)",
            "Phòng #R-1 đang ở trạng thái 'Phòng Đang Trống' (status=1)",
            "1. Tại danh sách phòng Admin, bấm dropdown 'Đổi Trạng Thái'\n2. Chọn '4. Đã Được Đặt (Booked)'",
            "Room ID: 1, Val: 4",
            "Trạng thái phòng chuyển sang Badge xanh dương 'Đã Được Đặt' (status=4). Cập nhật DB thành công.",
            "PASS"
        ),
        (
            "TC_STATE_TRANS_02",
            "Chuyển đổi trạng thái (State Transition)",
            "Chuyển trạng thái: Đã Được Đặt (status=4) -> Đang Có Khách (status=3)",
            "Chuyển đổi trạng thái (State Transition)",
            "Phòng #R-1 đang ở trạng thái 'Đã Được Đặt' (status=4)",
            "1. Khi khách đến Check-in, Admin/Lễ tân chọn '3. Đang Có Khách (Occupied)'",
            "Room ID: 1, Val: 3",
            "Trạng thái phòng chuyển sang Badge đỏ 'Đang Có Khách' (status=3). Cập nhật DB thành công.",
            "PASS"
        ),
        (
            "TC_STATE_TRANS_03",
            "Chuyển đổi trạng thái (State Transition)",
            "Chuyển trạng thái: Đang Có Khách (status=3) -> Đang Dọn Dẹp (status=2)",
            "Chuyển đổi trạng thái (State Transition)",
            "Phòng #R-1 đang ở trạng thái 'Đang Có Khách' (status=3)",
            "1. Khi khách Check-out trả phòng, Admin/Lễ tân chọn '2. Đang Dọn Dẹp (Cleaning)'",
            "Room ID: 1, Val: 2",
            "Trạng thái phòng chuyển sang Badge vàng 'Đang Dọn Dẹp' (status=2). Cập nhật DB thành công.",
            "PASS"
        ),
        (
            "TC_STATE_TRANS_04",
            "Chuyển đổi trạng thái (State Transition)",
            "Chuyển trạng thái: Đang Dọn Dẹp (status=2) -> Phòng Trống (status=1)",
            "Chuyển đổi trạng thái (State Transition)",
            "Phòng #R-1 đang ở trạng thái 'Đang Dọn Dẹp' (status=2)",
            "1. Sau khi nhân viên buồng phòng dọn xong, Admin chọn '1. Phòng Đang Trống (Available)'",
            "Room ID: 1, Val: 1",
            "Trạng thái phòng hoàn tất 1 chu kỳ chuyển đổi, quay về Badge xanh 'Phòng Đang Trống' (status=1). Sẵn sàng cho khách đặt mới.",
            "PASS"
        ),
        (
            "TC_STATE_TRANS_05",
            "Chuyển đổi trạng thái (State Transition)",
            "Chuyển trạng thái sang Tạm Dừng / Bảo Trì (status=0) từ bất kỳ trạng thái nào",
            "Chuyển đổi trạng thái (State Transition)",
            "Phòng #R-1 bị hỏng điều hòa",
            "1. Admin bấm 'Đổi Trạng Thái'\n2. Chọn '0. Tạm Dừng / Bảo Trì'",
            "Room ID: 1, Val: 0",
            "Trạng thái phòng chuyển sang Badge xám 'Tạm Dừng' (status=0). Phòng ngừng kinh doanh tạm thời.",
            "PASS"
        ),
        (
            "TC_STATE_TRANS_06",
            "Chuyển đổi trạng thái (State Transition)",
            "Khôi phục trạng thái từ Bảo Trì (status=0) -> Phòng Trống (status=1)",
            "Chuyển đổi trạng thái (State Transition)",
            "Phòng #R-1 đã sửa xong thiết bị (status=0)",
            "1. Admin chọn '1. Phòng Đang Trống (Available)'",
            "Room ID: 1, Val: 1",
            "Phòng hoạt động trở lại bình thường.",
            "PASS"
        ),

        # --- NHÓM 7: KIỂM THỬ GIAO DIỆN WEB CLIENT & SECURITY GUARD URL ---
        (
            "TC_CLIENT_SEC_01",
            "Kiểm thử Web Client & Security",
            "Khách hàng tìm kiếm phòng trên Web: Phòng ở trạng thái 'Đang Dọn Dẹp' (status=2) KHÔNG HIỂN THỊ",
            "Kiểm thử nghiệp vụ / Lọc dữ liệu (Business Logic)",
            "Phòng #R-10 có status=2 (Đang Dọn Dẹp)",
            "1. Khách hàng mở trang `rooms.php` trên Web Client\n2. Chọn ngày nhận/trả phòng và bấm Tìm kiếm",
            "Query: `ajax/rooms.php?fetch_rooms`",
            "Phòng #R-10 KHÔNG XUẤT HIỆN trong danh sách phòng trên Web Client. Nút 'Đặt phòng ngay' không thể bấm được.",
            "PASS"
        ),
        (
            "TC_CLIENT_SEC_02",
            "Kiểm thử Web Client & Security",
            "Khách hàng tìm kiếm phòng: Phòng ở trạng thái 'Bảo Trì' (status=0) KHÔNG HIỂN THỊ",
            "Kiểm thử nghiệp vụ / Lọc dữ liệu (Business Logic)",
            "Phòng #R-11 có status=0 (Bảo trì)",
            "1. Khách hàng thực hiện tìm kiếm phòng trên Web",
            "Query Client",
            "Phòng #R-11 KHÔNG HIỂN THỊ trên giao diện Web Client (Do câu SQL filter `WHERE status=1`).",
            "PASS"
        ),
        (
            "TC_CLIENT_SEC_03",
            "Kiểm thử Web Client & Security",
            "BẢO MẬT: Khách cố tình truy cập trực tiếp URL `confirm_booking.php?id=X` với phòng đang 'Đang Dọn Dẹp' (status=2)",
            "Kiểm thử bảo mật truy cập (Security Access Control)",
            "Phòng ID #R-10 đang có status=2 (Đang Dọn Dẹp). Khách đã đăng nhập tài khoản.",
            "1. Khách nhập trực tiếp URL trên trình duyệt: `http://localhost/hotelbooking/confirm_booking.php?id=10`\n2. Nhấn Enter",
            "URL direct access: `id=10` (status=2)",
            "Hệ thống chạy Security Guard (`mysqli_num_rows($room_res) == 0`) và LẬP TỨC REDIRECT người dùng về trang `rooms.php`. Ngăn chặn tuyệt đối việc đặt phòng đang dọn dẹp!",
            "PASS"
        ),
        (
            "TC_CLIENT_SEC_04",
            "Kiểm thử Web Client & Security",
            "BẢO MẬT: Khách truy cập trực tiếp URL `confirm_booking.php?id=X` với phòng 'Bảo Trì' (status=0)",
            "Kiểm thử bảo mật truy cập (Security Access Control)",
            "Phòng ID #R-11 đang bảo trì (status=0)",
            "1. Khách nhập URL: `confirm_booking.php?id=11`\n2. Nhấn Enter",
            "URL direct access: `id=11` (status=0)",
            "Hệ thống tự động redirect chuyển hướng về trang `rooms.php`. Không thể thực hiện xác nhận đặt phòng.",
            "PASS"
        )
    ]

    start_row = 5
    for idx, tc in enumerate(testcases, 1):
        curr_row = start_row + idx - 1
        ws1.row_dimensions[curr_row].height = 50
        
        # Row data mapping
        row_data = [
            idx,          # STT
            tc[0],        # TC ID
            tc[1],        # Module / Sub-module
            tc[2],        # Title / Description
            tc[3],        # Technique
            tc[4],        # Pre-conditions
            tc[5],        # Steps
            tc[6],        # Test Data
            tc[7],        # Expected Result
            tc[8]         # Status
        ]
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=curr_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = border_all
            
            # Alignments
            if col_idx in [1, 2, 5, 10]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
            # Specific styling for columns
            if col_idx == 2:
                cell.font = id_font
            elif col_idx == 3:
                cell.font = group_font
            elif col_idx == 10:
                cell.fill = pass_fill
                cell.font = pass_font

    # Set column widths for Sheet 1
    col_widths1 = {
        1: 6,   # STT
        2: 20,  # Mã TC
        3: 25,  # Phân nhóm
        4: 35,  # Tên TC
        5: 22,  # Kỹ thuật
        6: 28,  # Pre-condition
        7: 40,  # Test Steps
        8: 25,  # Test Data
        9: 42,  # Expected Result
        10: 12  # Status
    }
    
    for col_idx, width in col_widths1.items():
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = width

    # -------------------------------------------------------------
    # SHEET 2: STATE TRANSITION MATRIX
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Ma Trận State Transition")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.cell(row=1, column=1, value="MA TRẬN CHUYỂN ĐỔI TRẠNG THÁI PHÒNG (STATE TRANSITION MATRIX)").font = title_font
    ws2.cell(row=2, column=1, value="Quy trình 5 trạng thái phòng: 1-Trống, 4-Đã đặt, 3-Đang ở, 2-Dọn dẹp, 0-Bảo trì").font = sub_title_font
    
    matrix_headers = [
        "STT",
        "Trạng Thái Hiện Tại (From State)",
        "Hành Động / Thao Tác (Action / Event)",
        "Trạng Thái Tiếp Theo (To State)",
        "Hợp Lệ / Không Hợp Lệ",
        "Cơ Chế Backend & Phản Ứng Web Client (Behavior)"
    ]
    
    ws2.row_dimensions[4].height = 28
    for col_idx, h in enumerate(matrix_headers, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = header_border

    matrix_rows = [
        (1, "S1: Phòng Trống (status=1)", "Khách đặt thành công trên Web / Admin gán đặt", "S4: Đã Được Đặt (status=4)", "HỢP LỆ", "Tạo bản ghi đơn đặt phòng mới trong `booking_order`, cập nhật `rooms.status = 4`."),
        (2, "S4: Đã Được Đặt (status=4)", "Lễ tân thực hiện Check-in nhận phòng", "S3: Đang Có Khách (status=3)", "HỢP LỆ", "Admin chuyển trạng thái sang 3. Khách lưu trú chính thức tại phòng."),
        (3, "S3: Đang Có Khách (status=3)", "Lễ tân thực hiện Check-out trả phòng", "S2: Đang Dọn Dẹp (status=2)", "HỢP LỆ", "Khách trả phòng. Admin chuyển phòng sang trạng thái 2 (Vàng) để buồng phòng vào dọn."),
        (4, "S2: Đang Dọn Dẹp (status=2)", "Nhân viên buồng phòng dọn dẹp xong", "S1: Phòng Trống (status=1)", "HỢP LỆ", "Admin bấm chuyển trạng thái 1 (Xanh). Phòng quay lại vòng lặp sẵn sàng đón khách."),
        (5, "S2: Đang Dọn Dẹp (status=2)", "Khách hàng truy cập Web tìm & đặt phòng", "KHÔNG THỂ ĐẶT PHÒNG", "BẢO VỆ AN TOÀN", "SQL `ajax/rooms.php` filter `status=1` loại bỏ phòng. Nút đặt không xuất hiện. Nếu gõ URL direct `confirm_booking.php?id=X` bị Security Guard redirect ngay về `rooms.php`."),
        (6, "S0: Bảo Trì (status=0)", "Khách hàng truy cập Web tìm & đặt phòng", "KHÔNG THỂ ĐẶT PHÒNG", "BẢO VỆ AN TOÀN", "Phòng bị ẩn hoàn toàn khỏi danh sách tìm kiếm và bị chặn bởi URL Direct Guard."),
        (7, "S1: Phòng Trống (status=1)", "Phát hiện sự cố hỏng hóc / Đưa đi bảo trì", "S0: Bảo Trì (status=0)", "HỢP LỆ", "Admin chọn trạng thái 0 (Xám). Tạm thời ngừng kinh doanh phòng."),
        (8, "S0: Bảo Trì (status=0)", "Hoàn tất sửa chữa & Dọn dẹp", "S1: Phòng Trống (status=1)", "HỢP LỆ", "Admin kích hoạt lại phòng về trạng thái 1.")
    ]

    for row_idx, r_data in enumerate(matrix_rows, 5):
        ws2.row_dimensions[row_idx].height = 42
        for col_idx, val in enumerate(r_data, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = border_all
            if col_idx in [1, 5]:
                cell.alignment = align_center
                if str(val) == "HỢP LỆ":
                    cell.font = Font(name=font_family, size=9.5, bold=True, color="16A34A")
                elif "BẢO VỆ" in str(val):
                    cell.font = Font(name=font_family, size=9.5, bold=True, color="0284C7")
            else:
                cell.alignment = align_left

    col_widths2 = {1: 6, 2: 28, 3: 35, 4: 28, 5: 18, 6: 50}
    for col_idx, width in col_widths2.items():
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = width

    # Save workbook
    output_filename = "TestCases_QuanLyPhong_StateTransition.xlsx"
    target_path = os.path.join(r"e:\School\KiemThu\Hotel-Booking-Website", output_filename)
    wb.save(target_path)
    print(f"Successfully generated Excel file at: {target_path}")

if __name__ == "__main__":
    generate_excel()
